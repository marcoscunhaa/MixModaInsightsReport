"""Export stock alerts to a spreadsheet organized by store."""
from collections import defaultdict
from pathlib import Path
import re

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from utils.formatting import CURRENCY_SYMBOL


NAVY = "14284B"
BLUE = "2854D9"
PALE = "F3F6FB"
WHITE = "FFFFFF"
INK = "233451"
MUTED = "718097"
GREEN = "E7F5EE"
GREEN_INK = "18794E"
LINE = "DCE4EF"
CHECK_PENDING = "☐ Pending"
CHECK_DONE = "☑ Picked"


def _sheet_name(store, used):
    name = re.sub(r"[\\/*?:\[\]]", " ", str(store)).strip().strip("'")
    name = re.sub(r"\s+", " ", name)[:31] or "Store"
    candidate = name
    suffix = 2
    while candidate.casefold() in used:
        ending = f" ({suffix})"
        candidate = f"{name[:31 - len(ending)]}{ending}"
        suffix += 1
    used.add(candidate.casefold())
    return candidate


def export_stock_alerts_xlsx(path, alerts):
    """Write one sortable, checkable worksheet per store with active alerts."""
    grouped = defaultdict(list)
    for alert in alerts:
        grouped[str(alert.get("loja") or "Unnamed store")].append(alert)

    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)
    used_names = set()

    for store in sorted(grouped, key=str.casefold):
        sheet = workbook.create_sheet(_sheet_name(store, used_names))
        sheet.sheet_view.showGridLines = False
        sheet.merge_cells("A1:G1")
        sheet["A1"] = f"STOCK REPLENISHMENT  •  {store}"
        sheet["A1"].font = Font(name="Aptos Display", size=16, bold=True, color=WHITE)
        sheet["A1"].fill = PatternFill("solid", fgColor=NAVY)
        sheet["A1"].alignment = Alignment(vertical="center")
        sheet.row_dimensions[1].height = 32

        sheet.merge_cells("A2:G2")
        sheet["A2"] = "Set the first column to “☑ Picked” once the goods have been picked."
        sheet["A2"].font = Font(name="Aptos", size=10, color=MUTED, italic=True)
        sheet["A2"].alignment = Alignment(vertical="center")
        sheet.row_dimensions[2].height = 24

        headers = [
            "Picked?", "Barcode", "Product", "Store",
            "Sale price", "Current stock", "Minimum stock",
        ]
        header_row = 4
        for column, title in enumerate(headers, 1):
            cell = sheet.cell(header_row, column, title)
            cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
            cell.fill = PatternFill("solid", fgColor=BLUE)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(bottom=Side(style="medium", color=NAVY))
        sheet.row_dimensions[header_row].height = 26

        rows = sorted(
            grouped[store],
            key=lambda item: (
                float(item.get("quantidade_estoque", 0) or 0),
                str(item.get("descricao", "")).casefold(),
            ),
        )
        for row_index, alert in enumerate(rows, header_row + 1):
            values = [
                CHECK_PENDING,
                str(alert.get("codigo_barras", "")),
                str(alert.get("descricao", "")),
                store,
                float(alert.get("preco_venda", 0) or 0),
                float(alert.get("quantidade_estoque", 0) or 0),
                int(alert.get("estoque_minimo", 0) or 0),
            ]
            for column, value in enumerate(values, 1):
                cell = sheet.cell(row_index, column, value)
                cell.font = Font(name="Aptos", size=10, color=INK)
                cell.fill = PatternFill("solid", fgColor=WHITE if row_index % 2 else PALE)
                cell.alignment = Alignment(
                    horizontal="center" if column in (1, 5, 6, 7) else "left",
                    vertical="center",
                )
                cell.border = Border(bottom=Side(style="hair", color=LINE))
            sheet.cell(row_index, 1).alignment = Alignment(horizontal="center", vertical="center")
            sheet.cell(row_index, 1).font = Font(name="Segoe UI Symbol", size=10, bold=True, color=BLUE)
            sheet.cell(row_index, 5).number_format = f'"{CURRENCY_SYMBOL}" #,##0.00'
            sheet.cell(row_index, 6).number_format = "0.##"
            sheet.cell(row_index, 7).number_format = "0"
            sheet.row_dimensions[row_index].height = 23

        last_row = header_row + len(rows)
        if rows:
            sheet.auto_filter.ref = f"A{header_row}:G{last_row}"
            sheet.freeze_panes = "A5"
            validation = DataValidation(
                type="list",
                formula1=f'"{CHECK_PENDING},{CHECK_DONE}"',
                allow_blank=False,
            )
            validation.promptTitle = "Product picking"
            validation.prompt = "Select Pending or Picked."
            validation.errorTitle = "Invalid option"
            validation.error = "Use one of the options in the list."
            validation.showErrorMessage = True
            validation.showInputMessage = True
            sheet.add_data_validation(validation)
            validation.add(f"A5:A{last_row}")
            sheet.conditional_formatting.add(
                f"A5:A{last_row}",
                FormulaRule(
                    formula=[f'$A5="{CHECK_DONE}"'],
                    fill=PatternFill("solid", fgColor=GREEN),
                    font=Font(color=GREEN_INK, bold=True),
                ),
            )

        for column, width in {
            "A": 19, "B": 22, "C": 48, "D": 22, "E": 18, "F": 16, "G": 18,
        }.items():
            sheet.column_dimensions[column].width = width
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = "1:4"
        sheet.sheet_view.zoomScale = 90

    if not grouped:
        sheet = workbook.create_sheet("No alerts")
        sheet["A1"] = "No product reached the stock rules."
        sheet["A1"].font = Font(name="Aptos", size=12, color=INK)

    workbook.save(Path(path))
